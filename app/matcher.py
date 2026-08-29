"""Поиск соответствия артикул -> файл PDF, парсинг входных списков артикулов."""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

ARTICLE_RE = re.compile(r'^[A-Za-zА-Яа-я0-9]+(?:-[A-Za-zА-Яа-я0-9]+)+$')


def _token_pattern(code: str) -> re.Pattern:
    return re.compile(
        r'(?:(?<=^)|(?<=[\s,.\(\)_]))' + re.escape(code) + r'(?=[,\s.\(\)_]|$)'
    )


def find_exact(code: str, files: list[str]) -> list[str]:
    pat = _token_pattern(code)
    return [f for f in files if pat.search(f)]


def find_with_fallbacks(code: str, files: list[str]) -> tuple[list[str], str]:
    """Возвращает (кандидаты, метод_совпадения)."""
    candidates = find_exact(code, files)
    if candidates:
        return candidates, "exact"

    # Фолбэк 2: код без размера + отдельное поле "размер X"
    if "-" in code:
        base, size = code.rsplit("-", 1)
        base_pat = _token_pattern(base)
        size_pat = re.compile(
            r'(?:размер\s*' + re.escape(size) + r'(?=[_\-,.\s]|$))'
            r'|(?:[_\-]' + re.escape(size) + r'(?=\.pdf$|[_\-]))',
            re.IGNORECASE,
        )
        candidates = [f for f in files if base_pat.search(f) and size_pat.search(f)]
        if candidates:
            return candidates, "fallback_size_field"

    # Фолбэк 3: кириллическая "М" вместо латинской в размере
    if code.endswith("-M") or code.endswith("-M".upper()):
        alt_code = code[:-1] + "М"  # заменить последнюю латинскую M на кириллическую М (U+041C)
        candidates = find_exact(alt_code, files)
        if candidates:
            return candidates, "fallback_cyrillic_m"

    return [], "not_found"


@dataclass
class ParsedList:
    codes: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def parse_text_list(text: str) -> ParsedList:
    result = ParsedList()
    for raw_line in text.splitlines():
        for token in raw_line.replace(",", " ").split():
            token = token.strip()
            if token:
                result.codes.append(token)
    return result


def parse_csv_ozon(path: str | Path) -> ParsedList:
    result = ParsedList()
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=";")
        if reader.fieldnames is None or "Артикул" not in reader.fieldnames:
            result.warnings.append(
                f"Колонка 'Артикул' не найдена в CSV. Найдены колонки: {reader.fieldnames}"
            )
            return result
        for row in reader:
            val = (row.get("Артикул") or "").strip()
            if val:
                result.codes.append(val)
    return result


def parse_csv_lamoda(path: str | Path) -> ParsedList:
    result = ParsedList()
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=";")
        if reader.fieldnames is None or "Артикул товара" not in reader.fieldnames:
            result.warnings.append(
                f"Колонка 'Артикул товара' не найдена в CSV. Найдены колонки: {reader.fieldnames}"
            )
            return result
        for row in reader:
            val = (row.get("Артикул товара") or "").strip()
            if val:
                result.codes.append(val)
    return result


def parse_xlsx_yandex(path: str | Path) -> ParsedList:
    result = ParsedList()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet_name = "Список заказов"
    if sheet_name not in wb.sheetnames:
        result.warnings.append(
            f"Лист '{sheet_name}' не найден. Найдены листы: {wb.sheetnames}"
        )
        return result
    ws = wb[sheet_name]
    for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if len(row) < 4:
            continue
        val = row[3]  # колонка D
        if val is None:
            continue
        val = str(val).strip()
        if val:
            result.codes.append(val)
    return result


def parse_any_file(path: str | Path) -> ParsedList:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".xlsx":
        return parse_xlsx_yandex(p)
    if suffix == ".csv":
        # Пытаемся определить источник по заголовку
        with open(p, newline="", encoding="utf-8-sig") as f:
            header = f.readline()
        if "Артикул товара" in header:
            return parse_csv_lamoda(p)
        return parse_csv_ozon(p)
    result = ParsedList()
    result.warnings.append(f"Неизвестный формат файла: {suffix}")
    return result


def list_pdf_files(folder: str | Path) -> list[str]:
    p = Path(folder)
    return [f.name for f in p.iterdir() if f.is_file() and f.suffix.lower() == ".pdf"]
