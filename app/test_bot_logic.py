"""Тесты логики Telegram-бота: парсинг сообщений, резолв папки, retry-чтение, сборка."""
import os
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from pypdf import PdfReader, PdfWriter

from message_parser import parse_message
from pdf_builder import build_from_source, build_pdf, make_output_name
from pdf_io import PdfUnreadable, read_pdf_with_retry


def make_pdf(path: Path, n_pages: int):
    w = PdfWriter()
    for _ in range(n_pages):
        w.add_blank_page(width=200, height=300)
    with open(path, "wb") as f:
        w.write(f)


def test_parse_real_message():
    text = """RUSH114B-XXL
RUSH083B-XXL
BR-VEST001B-XXL

ozn

RUSH079B-L
V-SHORT001M-L

yand

папка br 28
"""
    m = parse_message(text)
    assert m.ozon == ["RUSH114B-XXL", "RUSH083B-XXL", "BR-VEST001B-XXL"], m.ozon
    assert m.yandex == ["RUSH079B-L", "V-SHORT001M-L"], m.yandex
    assert m.folder_query == "br 28" and not m.strict_folder
    print("OK: разбор реального сообщения (маркер ПОСЛЕ списка)")


def test_parse_marker_before_and_reverse_order():
    m = parse_message("YANDEX\nA-1\nB-2\nOZON\nC-3\nпапка FIZ 28")
    assert m.yandex == ["A-1", "B-2"], m.yandex
    assert m.ozon == ["C-3"], m.ozon
    assert m.folder_query == "FIZ 28"
    print("OK: маркер ПЕРЕД списком и обратный порядок блоков")


def test_parse_strict_folder_and_variants():
    for text, exp in [
        ("A-1\nozon\nbr 29 папка и только она", True),
        ("A-1\nozon\nпапка BR 29", False),
        ("A-1\nozon\nтолько папка BR 29", True),
    ]:
        m = parse_message(text)
        assert m.folder_query.upper().replace(" ", "") == "BR29", m.folder_query
        assert m.strict_folder is exp, (text, m.strict_folder)
    print("OK: флаг «только эта папка» и варианты записи папки")


def test_parse_single_list_and_repeats():
    m = parse_message("SHORT001F-M\nSHORT001F-M\nозон\nпапка BR 29")
    assert m.ozon == ["SHORT001F-M", "SHORT001F-M"] and m.yandex == []
    assert not m.warnings
    m2 = parse_message("SHORT001F-M SHORT002F-L\nпапка BR 29")
    assert m2.ozon == ["SHORT001F-M", "SHORT002F-L"] and m2.warnings, m2
    print("OK: один список, повторы сохранены, отсутствие маркера отмечено предупреждением")


def test_folder_resolution():
    tmp = Path(tempfile.mkdtemp(prefix="cz_roots_"))
    (tmp / "BR 29").mkdir()
    (tmp / "BR 291").mkdir()
    (tmp / "FIZ OZN").mkdir()
    make_pdf(tmp / "BR 29" / "x_A-1, размер M_1.pdf", 1)
    os.environ["CZ_LABEL_ROOTS"] = str(tmp)
    import importlib

    import folders as folders_mod
    importlib.reload(folders_mod)

    assert folders_mod.resolve_folder("br 29").match == tmp / "BR 29"
    assert folders_mod.resolve_folder("fiz ozn").match == tmp / "FIZ OZN"
    assert folders_mod.resolve_folder("несуществующая").status == "not_found"
    assert folders_mod.resolve_folder(None).status == "no_query"
    # BR 2 -> два кандидата по префиксу, но PDF только в одном -> выбираем его
    assert folders_mod.resolve_folder("BR 2").match == tmp / "BR 29"
    (tmp / "BR 291" / "y.pdf").write_bytes((tmp / "BR 29" / "x_A-1, размер M_1.pdf").read_bytes())
    importlib.reload(folders_mod)
    assert folders_mod.resolve_folder("BR 2").status == "ambiguous"
    print("OK: резолв папки (точное/префикс/неоднозначность/не найдено)")
    shutil.rmtree(tmp, ignore_errors=True)
    os.environ.pop("CZ_LABEL_ROOTS", None)


def test_read_retry_and_unreadable():
    tmp = Path(tempfile.mkdtemp(prefix="cz_read_"))
    good = tmp / "04630354081712_Шорты, FIZULI, A-CODE001F-M, цвет черный, размер M_2.pdf"
    make_pdf(good, 2)
    assert len(read_pdf_with_retry(good).pages) == 2

    bad = tmp / "04630354081713_Шорты, FIZULI, B-CODE002F-L, цвет синий, размер L_1.pdf"
    bad.write_bytes(b"not a pdf at all")
    try:
        read_pdf_with_retry(bad, retries=2, delay=0.01)
        raise AssertionError("ожидалось PdfUnreadable")
    except PdfUnreadable:
        pass

    report = build_pdf(["A-CODE001F-M", "B-CODE002F-L", "NOPE-X-L"], tmp, tmp / "out.pdf")
    assert report.found == 1, report
    assert report.unreadable_codes == ["B-CODE002F-L"], report.unreadable_codes
    assert report.missing_codes == ["NOPE-X-L"], report.missing_codes
    assert len(PdfReader(str(tmp / "out.pdf")).pages) == 1
    print("OK: retry-чтение, «не прочитан» отделён от «не найден»")
    shutil.rmtree(tmp, ignore_errors=True)


def test_end_to_end_counts():
    """Списки Ozon+Yandex, фолбэк по размеру, повторы, отсутствующий артикул."""
    tmp = Path(tempfile.mkdtemp(prefix="cz_e2e_"))
    make_pdf(tmp / "1_Худи, BARRACUDA, RUSH114B-XXL, цвет черный, размер XXL_1.pdf", 1)
    make_pdf(tmp / "2_Худи, BARRACUDA, RUSH083B-XXL, цвет серый, размер XXL_1.pdf", 1)
    make_pdf(tmp / "3_Жилет, BARRACUDA, BR-VEST001B-XXL, цвет хаки, размер XXL_1.pdf", 1)
    make_pdf(tmp / "4_Худи, BARRACUDA, RUSH079B-L, цвет синий, размер L_1.pdf", 1)
    # фолбэк: код без суффикса размера, размер отдельным полем
    make_pdf(tmp / "5_Шорты, BARRACUDA, V-SHORT001M, цвет черный, размер L_1.pdf", 1)
    # 3 страницы на артикул, а в списке он встречается 2 раза -> должно попасть 2 страницы
    make_pdf(tmp / "6_Футболка, FIZULI, SHORT001F-M, цвет белый, размер M_3.pdf", 3)

    text = """RUSH114B-XXL
RUSH083B-XXL
BR-VEST001B-XXL
SHORT001F-M
SHORT001F-M
SOCKS-FIZ001

ozn

RUSH079B-L
V-SHORT001M-L

yand

папка BR 28
"""
    m = parse_message(text)
    out_o = tmp / make_output_name("OZON", "BR28")
    rep_o = build_pdf(m.ozon, tmp, out_o)
    assert rep_o.total == 6 and rep_o.found == 5, rep_o
    assert rep_o.missing_codes == ["SOCKS-FIZ001"], rep_o.missing_codes
    assert len(PdfReader(str(out_o)).pages) == 5
    assert out_o.name == "OZON_BR28_CZ_po_spisku.pdf"

    out_y = tmp / make_output_name("YANDEX", "BR28")
    rep_y = build_pdf(m.yandex, tmp, out_y)
    assert rep_y.total == 2 and rep_y.found == 2, rep_y
    assert rep_y.fallback_used.get("V-SHORT001M-L") == "fallback_size_field", rep_y.fallback_used
    assert len(PdfReader(str(out_y)).pages) == 2
    print("OK: сквозной сценарий Ozon+Yandex (повторы, фолбэк, «не найдено», имена файлов)")
    shutil.rmtree(tmp, ignore_errors=True)


def test_pool_roundtrip_and_build():
    """Пул: склейка -> индекс внутри файла -> сборка по списку так же, как из папки."""
    import zipfile

    from pdf_builder import FolderSource, build_from_source
    from pool import build_pool, collect_from_zip, load_pool

    tmp = Path(tempfile.mkdtemp(prefix="cz_pool_"))
    src = tmp / "labels"
    src.mkdir()
    make_pdf(src / "1_Худи, BARRACUDA, RUSH114B-XXL, цвет черный, размер XXL_1.pdf", 1)
    make_pdf(src / "2_Футболка, FIZULI, SHORT001F-M, цвет белый, размер M_3.pdf", 3)
    make_pdf(src / "3_Шорты, BARRACUDA, V-SHORT001M, цвет черный, размер L_1.pdf", 1)
    sources = sorted((f.name, f) for f in src.iterdir())

    pool, unreadable = build_pool(sources, tmp / "pool.pdf", "BR 28")
    assert not unreadable and pool.page_count == 5 and len(pool.filenames) == 3, pool.page_files

    loaded = load_pool(tmp / "pool.pdf")
    assert loaded is not None, "индекс должен читаться из самого PDF, без внешних файлов"
    assert loaded.page_files == pool.page_files

    codes = ["RUSH114B-XXL", "SHORT001F-M", "SHORT001F-M", "V-SHORT001M-L", "SOCKS-FIZ001"]
    rep_pool = build_from_source(codes, loaded, tmp / "from_pool.pdf")
    rep_folder = build_from_source(codes, FolderSource(src), tmp / "from_folder.pdf")
    assert rep_pool.found == rep_folder.found == 4, (rep_pool, rep_folder)
    assert rep_pool.missing_codes == ["SOCKS-FIZ001"]
    assert rep_pool.fallback_used.get("V-SHORT001M-L") == "fallback_size_field"
    assert len(PdfReader(str(tmp / "from_pool.pdf")).pages) == 4
    print("OK: пул самодостаточен, результат совпадает со сборкой из папки")

    # обычный PDF без индекса пулом не считается
    make_pdf(tmp / "plain.pdf", 2)
    assert load_pool(tmp / "plain.pdf") is None
    print("OK: PDF без индекса не принимается за пул")

    # zip с кириллицей в именах и мусором __MACOSX
    zpath = tmp / "BR 28.zip"
    with zipfile.ZipFile(zpath, "w") as zf:
        for name, path in sources:
            zf.write(path, arcname=f"BR 28/{name}")
        zf.writestr("__MACOSX/._junk.pdf", b"junk")
        zf.writestr("readme.txt", b"not a pdf")
    got = collect_from_zip(zpath, tmp / "unzip")
    assert [n for n, _ in got] == [n for n, _ in sources], got
    pool2, _ = build_pool(got, tmp / "pool2.pdf", "BR 28")
    assert pool2.page_count == 5
    print("OK: распаковка zip (вложенные папки, кириллица, мусор __MACOSX)")
    shutil.rmtree(tmp, ignore_errors=True)


def test_archive_via_external_tool():
    """Не-zip архивы (rar/7z/tar) распаковываются внешним распаковщиком.

    Настоящий .rar на машине не собрать (нет rar-компрессора), поэтому проверяем
    тот же самый путь на .7z, который умеет создавать bsdtar.
    """
    import shutil as sh
    import subprocess

    from pool import build_pool, collect_from_archive

    if not sh.which("bsdtar"):
        print("SKIP: bsdtar недоступен")
        return

    tmp = Path(tempfile.mkdtemp(prefix="cz_arch_"))
    src = tmp / "labels"
    src.mkdir()
    make_pdf(src / "1_Худи, BARRACUDA, RUSH114B-XXL, цвет черный, размер XXL_2.pdf", 2)
    make_pdf(src / "2_Шорты, FIZULI, SHORT001F-M, цвет белый, размер M_1.pdf", 1)
    arch = tmp / "labels.7z"
    subprocess.run(["bsdtar", "-c", "--format", "7zip", "-f", str(arch), "-C", str(src), "."],
                   check=True, capture_output=True)

    import unicodedata

    got = collect_from_archive(arch, tmp / "unpacked")
    expected = sorted(unicodedata.normalize("NFC", f.name) for f in src.iterdir())
    assert [n for n, _ in got] == expected, (got, expected)
    pool, bad = build_pool(got, tmp / "pool.pdf", "arch")
    assert not bad and pool.page_count == 3, (bad, pool.page_files)
    print("OK: не-zip архив распакован внешним распаковщиком (кириллица в именах сохранена)")
    shutil.rmtree(tmp, ignore_errors=True)


def test_caption_as_list_and_zip_hygiene():
    """Регрессия по реальному прогону: архив + список в подписи одним сообщением."""
    import zipfile

    from pool import build_pool, collect_from_zip
    from tg_bot import jobs_from, pool_name_from

    caption = "A-CODE001F-M\nA-CODE001F-M\nB-CODE002F-L\n\nozon\n\nC-CODE003F-S\n\nyandex"
    # подпись со списком не должна становиться именем пула
    assert pool_name_from(caption, "FIZ 29") == "FIZ 29"
    assert pool_name_from("BR 30", "FIZ 29") == "BR 30"
    jobs = jobs_from(parse_message(caption))
    assert [(j.marketplace, len(j.codes)) for j in jobs] == [("OZON", 3), ("YANDEX", 1)], jobs

    tmp = Path(tempfile.mkdtemp(prefix="cz_cap_"))
    src = tmp / "labels"
    src.mkdir()
    make_pdf(src / "1_Шорты, FIZULI, A-CODE001F-M, цвет черный, размер M_2.pdf", 2)
    make_pdf(src / "2_Жилет, FIZULI, B-CODE002F-L, цвет хаки, размер L_1.pdf", 1)
    make_pdf(src / "3_Майка, FIZULI, C-CODE003F-S, цвет белый, размер S_1.pdf", 1)
    make_pdf(src / "OZON_FIZ29_CZ_po_spisku.pdf", 20)  # результат прошлой сборки, случайно в архиве

    # Имя, записанное в UTF-8 без флага 0x800 (так делают многие архиваторы):
    # zipfile отдаёт его декодированным как cp437, из-за чего раньше выходили кракозябры.
    from pool import _decode_zip_name

    real_name = "1_Шорты, FIZULI, A-CODE001F-M, цвет черный, размер M_2.pdf"
    info = zipfile.ZipInfo(real_name.encode("utf-8").decode("cp437"))
    info.flag_bits &= ~0x800
    decoded = _decode_zip_name(info.filename, info)
    assert decoded == real_name, f"кракозябры: {decoded}"

    zpath = tmp / "FIZ 29.zip"
    with zipfile.ZipFile(zpath, "w") as zf:
        for f in sorted(src.iterdir()):
            zf.write(f, arcname=f"FIZ 29/{f.name}")
        zf.writestr("__MACOSX/._1.pdf", b"junk")

    got = collect_from_zip(zpath, tmp / "unzip")
    names = [n for n, _ in got]
    assert any("Шорты" in n for n in names), names
    assert not any("_CZ_po_spisku" in n for n in names), "результат прошлой сборки не должен попадать в пул"
    assert len(got) == 3, names

    pool, _ = build_pool(got, tmp / "pool.pdf", "FIZ 29")
    assert pool.page_count == 4, pool.page_files
    rep = build_from_source(jobs[0].codes, pool, tmp / "out.pdf")
    assert rep.found == 3 and len(PdfReader(str(tmp / "out.pdf")).pages) == 3, rep
    print("OK: список в подписи к архиву, UTF-8 без флага, результаты прошлых сборок отфильтрованы")
    shutil.rmtree(tmp, ignore_errors=True)


def test_summary_format():
    """Формат итога: «ozon 29/31, yandex 8/8.» + причины."""
    from tg_bot import Job, format_summary

    tmp = Path(tempfile.mkdtemp(prefix="cz_sum_"))
    make_pdf(tmp / "1_Худи, FIZULI, A-CODE001F-M, цвет черный, размер M_3.pdf", 3)
    make_pdf(tmp / "2_Майка, FIZULI, B-CODE002F-XL, цвет белый, размер XL_1.pdf", 1)

    ok = [(Job("OZON", ["A-CODE001F-M"]), build_pdf(["A-CODE001F-M"], tmp, tmp / "a.pdf")),
          (Job("YANDEX", ["B-CODE002F-XL"]), build_pdf(["B-CODE002F-XL"], tmp, tmp / "b.pdf"))]
    assert format_summary(ok, "FIZ 29") == "ozon 1/1, yandex 1/1.", format_summary(ok, "FIZ 29")

    codes = ["A-CODE001F-M", "SOCKS-FIZ001", "SOCKS-FIZ001"]
    bad = [(Job("OZON", codes), build_pdf(codes, tmp, tmp / "c.pdf"))]
    text = format_summary(bad, "FIZ 29")
    assert text.splitlines()[0] == "ozon 1/3."
    assert text.splitlines()[1] == (
        "не найдено в папке FIZ 29: SOCKS-FIZ001 (нужно 2 шт, файла нет вообще)."
    ), text

    # позиций больше, чем страниц в этикетке — обязательно предупредить
    short = ["B-CODE002F-XL", "B-CODE002F-XL"]
    rep = build_pdf(short, tmp, tmp / "d.pdf")
    assert rep.short_pages == {"B-CODE002F-XL": (2, 1)}, rep.short_pages
    text = format_summary([(Job("OZON", short), rep)], "FIZ 29")
    assert "не хватило страниц: B-CODE002F-XL (нужно 2, в файле 1)" in text, text
    print("OK: формат итога и предупреждение о нехватке страниц")
    shutil.rmtree(tmp, ignore_errors=True)


def test_order_file_formats():
    """Файл заказа: шаблон склада, шапки, количество, маркеры, csv."""
    import csv as _csv

    import openpyxl

    from order_file import make_template, parse_order_file

    tmp = Path(tempfile.mkdtemp(prefix="cz_order_"))

    # 1. Шаблон склада: без шапки, артикул в колонке A, название в B, хвост пустых строк
    p1 = tmp / "sklad.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None, "\xa0\xa0"])
    for code, name in (
        ("VDLZK004F-S", "Водолазка лапша"),
        ("FL-SHRT001H-XL", "Майка лапша"),
        ("FL-SHRT001H-XL", "Майка лапша"),
        ("G-LONG004F-XL", "Лонгслив"),
    ):
        ws.append([code, name])
    for _ in range(60):
        ws.append([None, None])
    wb.save(p1)
    r = parse_order_file(p1)
    assert r.lists == {"": ["VDLZK004F-S", "FL-SHRT001H-XL", "FL-SHRT001H-XL", "G-LONG004F-XL"]}, r.lists
    print("OK: шаблон склада — порядок и повторы сохранены, названия и пустые строки не мешают")

    # 2. Шапка + количество: 2 шт = две страницы
    p2 = tmp / "qty.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Артикул", "Кол-во"])
    ws.append(["A-CODE001F-M", 2])
    ws.append(["B-CODE002F-L", None])
    wb.save(p2)
    r = parse_order_file(p2)
    assert r.lists == {"": ["A-CODE001F-M", "A-CODE001F-M", "B-CODE002F-L"]}, r.lists
    assert "Кол-во" in r.note, r.note
    print("OK: колонка количества разворачивается в повторы")

    # 3. Маркеры ozon/yandex строкой внутри файла (как в текстовых сообщениях — маркер после списка)
    p3 = tmp / "markers.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["A-CODE001F-M", "назв"])
    ws.append(["ozon"])
    ws.append(["B-CODE002F-L", "назв"])
    ws.append(["yandex"])
    wb.save(p3)
    r = parse_order_file(p3)
    assert r.lists == {"OZON": ["A-CODE001F-M"], "YANDEX": ["B-CODE002F-L"]}, r.lists
    print("OK: строки-маркеры делят файл на списки Ozon/Yandex")

    # 4. CSV выгрузка Ozon
    p4 = tmp / "ozon.csv"
    with open(p4, "w", newline="", encoding="utf-8-sig") as f:
        w = _csv.writer(f, delimiter=";")
        w.writerow(["Номер отправления", "Артикул", "Статус"])
        w.writerow(["1", "A-CODE001F-M", "x"])
        w.writerow(["2", "A-CODE001F-M", "x"])
    r = parse_order_file(p4)
    assert r.lists == {"": ["A-CODE001F-M", "A-CODE001F-M"]}, r.lists
    print("OK: csv-выгрузка Ozon")

    # 5. Наш шаблон читается сам собой
    r = parse_order_file(make_template(tmp / "tpl.xlsx"))
    assert r.lists == {"OZON": ["RUSH114B-XXL", "SHORT001F-M", "SHORT001F-M"],
                       "YANDEX": ["V-SHORT001M-L"]}, r.lists
    print("OK: /template читается собственным разбором")

    # 6. Файл без артикулов — понятное предупреждение, без падения
    p6 = tmp / "junk.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["дата", "сумма"])
    ws.append(["01.09", 100])
    wb.save(p6)
    r = parse_order_file(p6)
    assert not r.total and r.warnings, r
    print("OK: файл без артикулов — предупреждение, не падение")
    shutil.rmtree(tmp, ignore_errors=True)


def run():
    test_parse_real_message()
    test_parse_marker_before_and_reverse_order()
    test_parse_strict_folder_and_variants()
    test_parse_single_list_and_repeats()
    test_folder_resolution()
    test_read_retry_and_unreadable()
    test_end_to_end_counts()
    test_pool_roundtrip_and_build()
    test_archive_via_external_tool()
    test_caption_as_list_and_zip_hygiene()
    test_summary_format()
    test_order_file_formats()
    print("\nВСЕ ТЕСТЫ БОТА ПРОЙДЕНЫ")


if __name__ == "__main__":
    run()
