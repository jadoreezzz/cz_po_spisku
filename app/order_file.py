"""Разбор файла заказа (xlsx / csv) в список артикулов.

В отличие от жёстко зашитых выгрузок маркетплейсов, здесь формат определяется
на месте: ищем колонку с артикулами, необязательную колонку количества
и необязательную колонку маркетплейса. Порядок строк сохраняется —
он и задаёт порядок страниц в готовом PDF.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

ARTICLE_RE = re.compile(r"^[A-Za-zА-Яа-я0-9]+(?:[-/][A-Za-zА-Яа-я0-9]+)+$")

ARTICLE_HEADERS = (
    "артикул", "артикул товара", "артикул продавца", "ваш sku", "sku",
    "shop sku", "offer id", "код товара", "article", "модель",
)
QTY_HEADERS = ("количество", "кол-во", "колво", "шт", "штук", "qty", "quantity", "count")
MARKET_HEADERS = ("маркетплейс", "площадка", "marketplace", "канал", "магазин")

OZON_WORDS = ("ozon", "озон", "ozn")
YANDEX_WORDS = ("yandex", "яндекс", "yand", "ym", "я.маркет")

MAX_QTY = 500  # защита от опечатки вида 10000 в колонке количества


@dataclass
class OrderParse:
    lists: dict[str, list[str]] = field(default_factory=dict)  # "OZON"/"YANDEX"/"" -> коды
    note: str = ""                                             # что распозналось
    warnings: list[str] = field(default_factory=list)

    @property
    def total(self) -> int:
        return sum(len(v) for v in self.lists.values())


def _norm(v) -> str:
    return str(v).strip().lower().replace("ё", "е") if v is not None else ""


def _header_base(v) -> str:
    """«Количество (необязательно)» -> «количество»: скобки и хвосты отбрасываем."""
    name = re.sub(r"\(.*?\)", " ", _norm(v))
    name = re.sub(r"[^0-9a-zа-я ]+", " ", name)
    return " ".join(name.split())


def _is_header(v, candidates: tuple[str, ...]) -> bool:
    base = _header_base(v)
    if not base:
        return False
    # кандидаты приводим к той же форме: «кол-во» -> «кол во»
    return any(base == c or base.startswith(c + " ")
               for c in (_header_base(x) for x in candidates))


def _looks_like_code(v) -> bool:
    return bool(v) and bool(ARTICLE_RE.match(str(v).strip()))


def _qty_of(v) -> int:
    if v is None or str(v).strip() == "":
        return 1
    try:
        n = int(float(str(v).strip().replace(",", ".")))
    except ValueError:
        return 1
    return max(1, min(n, MAX_QTY))


def _market_of(v) -> str:
    low = _norm(v)
    if any(w in low for w in OZON_WORDS):
        return "OZON"
    if any(w in low for w in YANDEX_WORDS):
        return "YANDEX"
    return ""


def _find_header(rows: list[list], limit: int = 15) -> tuple[int, dict[str, int]] | None:
    """Строка заголовков и индексы нужных колонок."""
    for i, row in enumerate(rows[:limit]):
        cols: dict[str, int] = {}
        for j, cell in enumerate(row):
            name = _norm(cell)
            if not name:
                continue
            if "article" not in cols and _is_header(cell, ARTICLE_HEADERS):
                cols["article"] = j
            elif "qty" not in cols and _is_header(cell, QTY_HEADERS):
                cols["qty"] = j
            elif "market" not in cols and _is_header(cell, MARKET_HEADERS):
                cols["market"] = j
        if "article" in cols:
            return i, cols
    return None


def _guess_article_column(rows: list[list]) -> int | None:
    """Заголовков нет — берём колонку, где больше всего похожего на артикул."""
    scores: dict[int, int] = {}
    for row in rows[:300]:
        for j, cell in enumerate(row):
            if _looks_like_code(cell):
                scores[j] = scores.get(j, 0) + 1
    if not scores:
        return None
    best = max(scores, key=lambda j: scores[j])
    return best if scores[best] >= 1 else None


def _rows_from_xlsx(path: Path) -> list[tuple[str, list[list]]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if any(any(c is not None and str(c).strip() for c in r) for r in rows):
            out.append((name, rows))
    return out


def _rows_from_csv(path: Path) -> list[tuple[str, list[list]]]:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = raw[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") >= sample.count(",") else ","
    rows = [list(r) for r in csv.reader(raw.splitlines(), delimiter=delim)]
    return [("csv", rows)]


def _marker_row(row: list) -> str | None:
    """Строка-разделитель вида «ozon» / «яндекс» внутри таблицы."""
    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
    if len(cells) != 1:
        return None
    return _market_of(cells[0]) or None


def _parse_rows(rows: list[list]) -> tuple[dict[str, list[str]], str] | None:
    header = _find_header(rows)
    if header is not None:
        start, cols = header
        a, q, m = cols["article"], cols.get("qty"), cols.get("market")
        detected = f"колонка артикулов «{str(rows[start][a]).strip()}»"
        if q is not None:
            detected += f", количество «{str(rows[start][q]).strip()}»"
        if m is not None:
            detected += f", маркетплейс «{str(rows[start][m]).strip()}»"
        data = rows[start + 1:]
    else:
        a = _guess_article_column(rows)
        if a is None:
            return None
        q = m = None
        detected = f"колонка {a + 1} (заголовков не нашёл, определил по содержимому)"
        data = rows

    lists: dict[str, list[str]] = {}
    pending: list[str] = []   # коды до строки-маркера (маркер часто идёт ПОСЛЕ списка)
    active = ""               # маркер, объявленный ПЕРЕД списком

    for row in data:
        marker = _marker_row(row) if m is None else None
        if marker:
            if pending:
                lists.setdefault(marker, []).extend(pending)
                pending = []
                active = ""
            else:
                active = marker
            continue

        if a >= len(row):
            continue
        code = row[a]
        if code is None or not str(code).strip():
            continue
        code = str(code).strip()
        if _is_header(code, ARTICLE_HEADERS):
            continue  # повторная шапка посреди файла
        if not _looks_like_code(code) and header is None:
            continue  # без шапки доверяем только тому, что похоже на артикул
        qty = _qty_of(row[q]) if q is not None and q < len(row) else 1

        if m is not None and m < len(row):
            lists.setdefault(_market_of(row[m]), []).extend([code] * qty)
        elif active:
            lists.setdefault(active, []).extend([code] * qty)
        else:
            pending.extend([code] * qty)

    if pending:
        lists.setdefault("", []).extend(pending)

    lists = {k: v for k, v in lists.items() if v}
    return (lists, detected) if lists else None


def parse_order_file(path: str | Path) -> OrderParse:
    p = Path(path)
    suffix = p.suffix.lower()
    result = OrderParse()

    try:
        sheets = _rows_from_xlsx(p) if suffix in (".xlsx", ".xlsm") else _rows_from_csv(p)
    except Exception as e:  # noqa: BLE001 — битый файл не должен ронять бота
        result.warnings.append(f"Не смог прочитать файл: {e}")
        return result

    best: tuple[dict[str, list[str]], str, str] | None = None
    for sheet_name, rows in sheets:
        parsed = _parse_rows(rows)
        if parsed is None:
            continue
        lists, detected = parsed
        total = sum(len(v) for v in lists.values())
        if best is None or total > sum(len(v) for v in best[0].values()):
            best = (lists, detected, sheet_name)

    if best is None:
        names = ", ".join(n for n, _ in sheets) or "—"
        result.warnings.append(
            "Не нашёл в файле колонку с артикулами. "
            f"Листы: {names}. Нужна колонка «Артикул» (или SKU) — "
            "по одной строке на позицию. Шаблон: /template"
        )
        return result

    lists, detected, sheet_name = best
    result.lists = lists
    sheet_note = f"лист «{sheet_name}», " if suffix in (".xlsx", ".xlsm") and len(sheets) > 1 else ""
    result.note = sheet_note + detected
    return result


def make_template(path: str | Path) -> Path:
    """Шаблон заказа в том же виде, в каком его заполняет склад:
    артикул в первой колонке, название — для глаз. Количество и маркетплейс
    необязательны: повтор позиции можно задать и просто второй строкой."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заказ"
    ws.append(["Артикул", "Название (необязательно)", "Количество (необязательно)",
               "Маркетплейс (необязательно)"])
    for row in (
        ["RUSH114B-XXL", "Худи BARRACUDA", 1, "ozon"],
        ["SHORT001F-M", "Шорты FIZULI", 2, "ozon"],
        ["V-SHORT001M-L", "Шорты FIZULI", 1, "yandex"],
    ):
        ws.append(row)
    for col, width in zip("ABCD", (26, 34, 26, 26)):
        ws.column_dimensions[col].width = width
    path = Path(path)
    wb.save(path)
    return path
